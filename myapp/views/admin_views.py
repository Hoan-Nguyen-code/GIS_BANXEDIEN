"""
Admin Views - WebGIS Xe Điện
Quản lý: Dashboard, Users, Kho, Tài chính, Đơn hàng, Trạm sạc, Thống kê
"""

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.db.models import Sum, Count, Avg
from django.utils import timezone
from datetime import timedelta
import json
from myapp.models import User  # ✅ Sửa chỗ 1: import đúng custom User
from myapp.models import Order
import openpyxl
from django.http import HttpResponse
from django.db.models import F
from myapp.models import StockIn, StockOut, Inventory, Order, Product
from openpyxl.styles import Font, Alignment, PatternFill
from django.core.paginator import Paginator
from django.contrib.auth.decorators import login_required, user_passes_test
# Decorator kiểm tra user là admin
def admin_required(user):
    return user.is_authenticated and user.role == User.Role.ADMIN  # ✅ Sửa chỗ 2

# ==================== DASHBOARD ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_dashboard(request):
    from myapp.models import Order, Product
    
    stats = {
        'total_users': User.objects.count(),
        'total_products': Product.objects.count(),
        'total_stations': 32,  # cập nhật sau khi có model trạm sạc
        'total_orders': Order.objects.count(),
        'revenue_month': 15000000000,
        'expense_month': 8000000000,
        'profit_month': 7000000000,
        'stock_low': 5,
    }

    recent_orders = Order.objects.select_related('user').prefetch_related('items__product').order_by('-created_at')[:5]

    context = {
        'stats': stats,
        'recent_orders': recent_orders,
        'admin_name': request.user.username,
    }

    return render(request, 'admin/admin_dashboard.html', context)

# ==================== QUẢN LÝ USERS ====================
from django.core.paginator import Paginator # Đảm bảo đã import cái này ở đầu file

@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_users(request):
    users_list = User.objects.all().order_by('-date_joined')
    
    # Chia 10 người dùng mỗi trang
    paginator = Paginator(users_list, 10) 
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    stats = {
        'total': users_list.count(),
        'active': users_list.filter(is_active=True).count(),
        'admin': users_list.filter(role=User.Role.ADMIN).count(),
        'new_this_month': users_list.filter(
            date_joined__gte=timezone.now() - timedelta(days=30)
        ).count()
    }
    
    context = {
        'page_obj': page_obj,  # Chuyển sang dùng page_obj để quản lý trang
        'stats': stats,
    }
    return render(request, 'admin/admin_users.html', context)


# ==================== QUẢN LÝ KHO ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_kho(request):
    # Xử lý khi nhấn nút Nhập hoặc Xuất trên giao diện
# Trong hàm admin_kho(request):
    if request.method == "POST":
        action = request.POST.get('action')
        product_id = request.POST.get('product_id')
        quantity = int(request.POST.get('quantity', 0))
        
        try:
            product = get_object_or_404(Product, id=product_id)
            if action == 'IMPORT':
                # Lấy giá nhập từ form, nếu không nhập thì lấy giá bán hiện tại làm mặc định
                import_price = request.POST.get('import_price') or product.price
                StockIn.objects.create(
                    product=product,
                    quantity=quantity,
                    import_price=import_price,
                    imported_by=request.user
                )
                messages.success(request, f"Đã nhập kho {quantity} {product.name}")
            
            elif action == 'EXPORT':
                # Model của Huy sẽ tự raise ValueError nếu không đủ hàng
                StockOut.objects.create(
                    product=product,
                    quantity=quantity,
                    exported_by=request.user
                )
                messages.success(request, f"Đã xuất kho {quantity} {product.name}")
                
        except ValueError as e:
            # Bắt lỗi "Không đủ hàng trong kho" từ Model
            messages.error(request, f"Lỗi: {str(e)}")
        except Exception as e:
            messages.error(request, f"Lỗi hệ thống: {str(e)}")
        
        return redirect('admin_kho')
# 2. XỬ LÝ HIỂN THỊ VÀ PHÂN TRANG (Phần Huy cần sửa)
    products_list = Product.objects.select_related('category', 'inventory').all().order_by('-id')
    
    # Lấy toàn bộ sản phẩm cho danh sách xổ xuống trong Modal
    all_products = Product.objects.all().order_by('name') 
    
    paginator = Paginator(products_list, 10) 
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    stats = {
        'total_products': products_list.count(),
        'in_stock': sum(1 for p in products_list if hasattr(p, 'inventory') and p.inventory.stock_quantity > 5),
        'low_stock': sum(1 for p in products_list if hasattr(p, 'inventory') and 0 < p.inventory.stock_quantity <= 5),
        'out_of_stock': sum(1 for p in products_list if not hasattr(p, 'inventory') or p.inventory.stock_quantity == 0),
    }

    # Trả về context đầy đủ
    return render(request, 'admin/admin_kho.html', {
        'page_obj': page_obj, 
        'all_products': all_products, # PHẢI CÓ DÒNG NÀY ĐỂ MODAL HIỆN DANH SÁCH
        'stats': stats
    })

    # Hiển thị dữ liệu như cũ
    products = Product.objects.select_related('category', 'inventory').all()
    stats = {
        'total_products': products.count(),
        'in_stock': sum(1 for p in products if hasattr(p, 'inventory') and p.inventory.stock_quantity > 5),
        'low_stock': sum(1 for p in products if hasattr(p, 'inventory') and 0 < p.inventory.stock_quantity <= 5),
        'out_of_stock': sum(1 for p in products if not hasattr(p, 'inventory') or p.inventory.stock_quantity == 0),
    }
    return render(request, 'admin/admin_kho.html', {'products': products, 'stats': stats})


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_product_detail(request, product_id):
    from myapp.models import Product
    product = get_object_or_404(Product, id=product_id)
    return render(request, 'admin/admin_product_detail.html', {'product': product})


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_product_add(request):
    from myapp.models import Product, Category, Inventory

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        category_id = request.POST.get('category')
        description = request.POST.get('description', '').strip()
        price = request.POST.get('price', '0')
        stock_quantity = int(request.POST.get('stock_quantity', 0))
        is_active = request.POST.get('is_active') == 'on'
        image = request.FILES.get('image')

        if not name or not category_id or not price:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin!')
            categories = Category.objects.all()
            return render(request, 'admin/admin_product_form.html', {
                'action': 'add', 'categories': categories
            })

        product = Product.objects.create(
            name=name,
            category_id=category_id,
            description=description,
            price=price,
            is_active=is_active,
            created_by=request.user,
            image=image if image else 'products/default.jpg',
        )

        Inventory.objects.create(
            product=product,
            stock_quantity=stock_quantity,
        )

        messages.success(request, f'Đã thêm sản phẩm "{name}" thành công!')
        return redirect('admin_kho')

    categories = Category.objects.all()
    return render(request, 'admin/admin_product_form.html', {
        'action': 'add',
        'categories': categories,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_product_edit(request, product_id):
    from myapp.models import Product, Category, Inventory

    product = get_object_or_404(Product, id=product_id)
    inventory, _ = Inventory.objects.get_or_create(product=product)

    if request.method == 'POST':
        # 1. Cập nhật thông tin sản phẩm
        product.name = request.POST.get('name', '').strip()
        product.category_id = request.POST.get('category')
        product.description = request.POST.get('description', '').strip()
        product.price = request.POST.get('price', product.price)
        product.is_active = request.POST.get('is_active') == 'on'
        if request.FILES.get('image'):
            product.image = request.FILES.get('image')
        product.save()

        # 2. Xử lý logic tồn kho (CHỖ NÀY NÈ)
        try:
            stock_input = int(request.POST.get('stock_quantity', inventory.stock_quantity))
            if stock_input < 0:
                messages.error(request, "Số lượng tồn kho không thể là số âm!")
                # Nếu lỗi, render lại form với dữ liệu hiện tại để user sửa lại
                categories = Category.objects.all()
                return render(request, 'admin/admin_product_form.html', {
                    'action': 'edit', 'product': product, 'categories': categories, 'inventory': inventory
                })
            
            inventory.stock_quantity = stock_input
            inventory.save()
            messages.success(request, f'Đã cập nhật sản phẩm "{product.name}" thành công!')
            return redirect('admin_kho')
            
        except ValueError:
            messages.error(request, "Vui lòng nhập số hợp lệ cho tồn kho!")

    # Nếu là GET request thì hiển thị form như bình thường
    categories = Category.objects.all()
    return render(request, 'admin/admin_product_form.html', {
        'action': 'edit',
        'product': product,
        'categories': categories,
        'inventory': inventory,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_product_delete(request, product_id):
    from myapp.models import Product
    product = get_object_or_404(Product, id=product_id)

    if request.method == 'POST':
        name = product.name
        product.delete()
        messages.success(request, f'Đã xóa sản phẩm "{name}" thành công!')
        return redirect('admin_kho')

    return render(request, 'admin/admin_product_confirm_delete.html', {'product': product})


# ==================== QUẢN LÝ TÀI CHÍNH ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_taichinh(request):
    from myapp.models import Order, OrderItem, Product
    from django.db.models import Sum, F, Count
    from django.utils import timezone
    import json

    now = timezone.now()
    
    # ✅ Lấy tháng/năm từ filter, mặc định là tháng hiện tại
    selected_month = int(request.GET.get('month', now.month))
    selected_year = int(request.GET.get('year', now.year))
    today = now.date()

    VALID_STATUS = ['CONFIRMED', 'COMPLETED']

    # Doanh thu tháng được chọn
    revenue_month = Order.objects.filter(
        created_at__month=selected_month,
        created_at__year=selected_year,
        status__in=VALID_STATUS
    ).aggregate(total=Sum('total_price'))['total'] or 0

    # Doanh thu hôm nay
    revenue_today = Order.objects.filter(
        created_at__date=today,
        status__in=VALID_STATUS
    ).aggregate(total=Sum('total_price'))['total'] or 0

    # Doanh thu theo tháng (12 tháng trong năm được chọn)
    monthly_revenue = []
    monthly_expense = []
    monthly_profit = []

    for month in range(1, 13):
        rev = Order.objects.filter(
            created_at__month=month,
            created_at__year=selected_year,
            status__in=VALID_STATUS
        ).aggregate(total=Sum('total_price'))['total'] or 0

        rev_billion = round(float(rev) / 1_000_000_000, 2)
        exp_billion = round(rev_billion * 0.6, 2)
        pro_billion = round(rev_billion * 0.4, 2)

        monthly_revenue.append(rev_billion)
        monthly_expense.append(exp_billion)
        monthly_profit.append(pro_billion)

    revenue_data = {
        'labels': ['T1','T2','T3','T4','T5','T6','T7','T8','T9','T10','T11','T12'],
        'revenue': monthly_revenue,
        'expense': monthly_expense,
        'profit': monthly_profit,
    }

    expense_month = round(float(revenue_month) * 0.6, 2)
    profit_month = round(float(revenue_month) * 0.4, 2)

    def fmt(num):
        return f"{int(num):,} VNĐ".replace(',', '.')

    stats = {
        'revenue_month': fmt(revenue_month),
        'expense_month': fmt(expense_month),
        'profit_month': fmt(profit_month),
        'revenue_today': fmt(revenue_today),
    }

    # Top sản phẩm
    top_products_qs = OrderItem.objects.values(
        'product__name'
    ).annotate(
        sold=Sum('quantity'),
        revenue=Sum(F('quantity') * F('price'))
    ).order_by('-revenue')[:5]

    top_products = []
    for p in top_products_qs:
        top_products.append({
            'name': p['product__name'],
            'sold': p['sold'],
            'revenue': f"{int(p['revenue']):,}".replace(',', '.'),
        })

    # ✅ Dữ liệu cho dropdown filter
    months = [{'value': m} for m in range(1, 13)]
    years = list(range(2024, now.year + 1))

    context = {
        'stats': stats,
        'revenue_data': json.dumps(revenue_data),
        'top_products': top_products,
        'selected_month': selected_month,
        'selected_year': selected_year,
        'months': months,
        'years': years,
    }

    return render(request, 'admin/admin_taichinh.html', context)

# ==================== QUẢN LÝ ĐƠN HÀNG ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_donhang(request):
    from myapp.models import Order
    from django.core.paginator import Paginator

    # 1. Lấy tham số lọc và trang hiện tại
    status_filter = request.GET.get('status', '')
    page_number = request.GET.get('page')

    # 2. Lấy danh sách gốc (tối ưu bằng select_related và prefetch_related)
    orders_queryset = Order.objects.all().select_related('user').prefetch_related('items__product')

    # 3. Áp dụng bộ lọc nếu có
    if status_filter:
        orders_queryset = orders_queryset.filter(status=status_filter)

    # Sắp xếp mới nhất lên đầu
    orders_queryset = orders_queryset.order_by('-created_at')

    # 4. THIẾT LẬP PHÂN TRANG (10 đơn hàng/trang)
    paginator = Paginator(orders_queryset, 10)
    page_obj = paginator.get_page(page_number)

    # 5. Thống kê (Tính trên toàn bộ database để con số luôn chính xác)
    stats = {
        'total': Order.objects.count(),
        'pending': Order.objects.filter(status='PENDING').count(),
        'confirmed': Order.objects.filter(status='CONFIRMED').count(),
        'shipping': Order.objects.filter(status='SHIPPED').count(),
        'completed': Order.objects.filter(status='COMPLETED').count(),
        'cancelled': Order.objects.filter(status='CANCELLED').count(),
    }

    # Trong admin_views.py
    context = {
        'page_obj': page_obj,   # ĐÚNG: page_obj
        'stats': stats,
        'status_filter': status_filter,
    }
    return render(request, 'admin/admin_donhang.html', context)

@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_order_detail(request, order_id):
    from myapp.models import Order
    order = get_object_or_404(Order, id=order_id)
    return render(request, 'admin/admin_order_detail.html', {'order': order})


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_order_edit(request, order_id):
    from myapp.models import Order
    order = get_object_or_404(Order, id=order_id)

    if request.method == 'POST':
        new_status = request.POST.get('status', order.status)
        order.status = new_status
        order.save()
        messages.success(request, f'Đã cập nhật trạng thái đơn hàng #{order.id}!')
        return redirect('admin_donhang')

    return render(request, 'admin/admin_order_edit.html', {
        'order': order,
        'status_choices': Order.OrderStatus.choices,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_order_delete(request, order_id):
    from myapp.models import Order
    order = get_object_or_404(Order, id=order_id)

    if request.method == 'POST':
        order.delete()
        messages.success(request, f'Đã xóa đơn hàng #{order_id}!')
        return redirect('admin_donhang')

    return render(request, 'admin/admin_order_confirm_delete.html', {'order': order})

# ==================== QUẢN LÝ TRẠM SẠC ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_tramsac(request):
    from myapp.models import ChargingStation
    from django.core.paginator import Paginator

    # 1. Lấy toàn bộ danh sách trạm sạc, sắp xếp theo ID mới nhất
    stations_list = ChargingStation.objects.all().order_by('-id')

    # 2. Thiết lập phân trang (10 trạm mỗi trang)
    paginator = Paginator(stations_list, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # 3. Thống kê (Tính trên stations_list để con số luôn chính xác)
    stats = {
        'total': stations_list.count(),
        'active': stations_list.filter(status='ACTIVE').count(),
        'maintenance': stations_list.filter(status='MAINTENANCE').count(),
        'inactive': stations_list.filter(status='INACTIVE').count(),
    }

    context = {
        'page_obj': page_obj,  # Huy nhớ đổi tên biến này nhé
        'stats': stats,
    }
    return render(request, 'admin/admin_tramsac.html', context)


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_station_detail(request, station_id):
    from myapp.models import ChargingStation
    station = get_object_or_404(ChargingStation, id=station_id)
    return render(request, 'admin/admin_station_detail.html', {'station': station})


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_station_add(request):
    from myapp.models import ChargingStation

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        address = request.POST.get('address', '').strip()
        latitude = request.POST.get('latitude', '0')
        longitude = request.POST.get('longitude', '0')
        charger_type = request.POST.get('charger_type', 'DC_FAST')
        power = request.POST.get('power', '').strip()
        total_ports = int(request.POST.get('total_ports', 0))
        available_ports = int(request.POST.get('available_ports', 0))
        status = request.POST.get('status', 'ACTIVE')
        image = request.FILES.get('image')

        if not name or not address:
            messages.error(request, 'Vui lòng điền đầy đủ thông tin!')
            return render(request, 'admin/admin_station_form.html', {
                'action': 'add',
                'charger_types': ChargingStation.ChargerType.choices,
                'status_choices': ChargingStation.StationStatus.choices,
            })

        station = ChargingStation.objects.create(
            name=name,
            address=address,
            latitude=float(latitude),
            longitude=float(longitude),
            charger_type=charger_type,
            power=power,
            total_ports=total_ports,
            available_ports=available_ports,
            status=status,
        )

        if image:
            station.image = image
            station.save()

        messages.success(request, f'Đã thêm trạm sạc "{name}" thành công!')
        return redirect('admin_tramsac')

    from myapp.models import ChargingStation
    return render(request, 'admin/admin_station_form.html', {
        'action': 'add',
        'charger_types': ChargingStation.ChargerType.choices,
        'status_choices': ChargingStation.StationStatus.choices,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_station_edit(request, station_id):
    from myapp.models import ChargingStation
    station = get_object_or_404(ChargingStation, id=station_id)

    if request.method == 'POST':
        station.name = request.POST.get('name', '').strip()
        station.address = request.POST.get('address', '').strip()
        station.latitude = float(request.POST.get('latitude', station.latitude))
        station.longitude = float(request.POST.get('longitude', station.longitude))
        station.charger_type = request.POST.get('charger_type', station.charger_type)
        station.power = request.POST.get('power', '').strip()
        station.total_ports = int(request.POST.get('total_ports', station.total_ports))
        station.available_ports = int(request.POST.get('available_ports', station.available_ports))
        station.status = request.POST.get('status', station.status)

        if request.FILES.get('image'):
            station.image = request.FILES.get('image')

        station.save()
        messages.success(request, f'Đã cập nhật trạm sạc "{station.name}" thành công!')
        return redirect('admin_tramsac')

    return render(request, 'admin/admin_station_form.html', {
        'action': 'edit',
        'station': station,
        'charger_types': ChargingStation.ChargerType.choices,
        'status_choices': ChargingStation.StationStatus.choices,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_station_delete(request, station_id):
    from myapp.models import ChargingStation
    station = get_object_or_404(ChargingStation, id=station_id)

    if request.method == 'POST':
        name = station.name
        station.delete()
        messages.success(request, f'Đã xóa trạm sạc "{name}" thành công!')
        return redirect('admin_tramsac')

    return render(request, 'admin/admin_station_confirm_delete.html', {'station': station})


# ==================== THỐNG KÊ & BÁO CÁO ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_thongke(request):
    from myapp.models import Order, OrderItem, Category
    from django.db.models import Sum, Count, F
    from django.utils import timezone
    import json

    now = timezone.now()
    this_year = now.year

    # Doanh thu theo 7 ngày gần nhất
    daily_labels = []
    daily_data = []
    for i in range(6, -1, -1):
        day = now - timezone.timedelta(days=i)
        label = ['CN','T2','T3','T4','T5','T6','T7'][day.weekday() % 7] if day.weekday() != 6 else 'CN'
        daily_labels.append(label)
        rev = Order.objects.filter(
            created_at__date=day.date(),
            status__in=['CONFIRMED', 'COMPLETED']
        ).aggregate(total=Sum('total_price'))['total'] or 0
        daily_data.append(round(float(rev) / 1_000_000, 2))

    # Phân bổ sản phẩm theo danh mục
    category_data = OrderItem.objects.values(
        'product__category__name'
    ).annotate(
        total=Sum('quantity')
    ).order_by('-total')

    cat_labels = [c['product__category__name'] for c in category_data] or ['Chưa có data']
    cat_data = [c['total'] for c in category_data] or [1]

    chart_data = {
        'daily_revenue': {
            'labels': daily_labels,
            'data': daily_data,
        },
        'product_distribution': {
            'labels': cat_labels,
            'data': cat_data,
        },
    }

    # Top khách hàng VIP
    top_customers_qs = Order.objects.values(
        'user__first_name', 'user__last_name', 'user__username'
    ).annotate(
        orders=Count('id'),
        spent=Sum('total_price')
    ).order_by('-spent')[:5]

    top_customers = []
    for c in top_customers_qs:
        full_name = f"{c['user__last_name']} {c['user__first_name']}".strip()
        if not full_name:
            full_name = c['user__username']
        top_customers.append({
            'name': full_name,
            'orders': c['orders'],
            'spent': f"{int(c['spent']):,}".replace(',', '.'),
        })

    context = {
        'chart_data': json.dumps(chart_data),
        'top_customers': top_customers,
    }

    return render(request, 'admin/admin_thongke.html', context)


# ==================== CRUD USERS ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_user_detail(request, user_id):
    user = get_object_or_404(User, id=user_id)
    return render(request, 'admin/admin_user_detail.html', {'user': user})


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_user_add(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '').strip()
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        phone = request.POST.get('phone', '').strip()
        address = request.POST.get('address', '').strip()
        role = request.POST.get('role', User.Role.CUSTOMER)
        is_active = request.POST.get('is_active') == 'on'

        if not username or not password:
            messages.error(request, 'Username và mật khẩu không được để trống!')
            return render(request, 'admin/admin_user_form.html', {
                'action': 'add', 'roles': User.Role.choices
            })

        if User.objects.filter(username=username).exists():
            messages.error(request, f'Username "{username}" đã tồn tại!')
            return render(request, 'admin/admin_user_form.html', {
                'action': 'add', 'roles': User.Role.choices
            })

        User.objects.create_user(
            username=username, email=email, password=password,
            first_name=first_name, last_name=last_name,
            phone=phone, address=address,
            role=role, is_active=is_active,
        )
        messages.success(request, f'Đã thêm user "{username}" thành công!')
        return redirect('admin_users')

    return render(request, 'admin/admin_user_form.html', {
        'action': 'add',
        'roles': User.Role.choices,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_user_edit(request, user_id):
    user = get_object_or_404(User, id=user_id)

    if request.method == 'POST':
        user.email = request.POST.get('email', '').strip()
        user.first_name = request.POST.get('first_name', '').strip()
        user.last_name = request.POST.get('last_name', '').strip()
        user.phone = request.POST.get('phone', '').strip()
        user.address = request.POST.get('address', '').strip()
        user.role = request.POST.get('role', user.role)
        user.is_active = request.POST.get('is_active') == 'on'

        new_password = request.POST.get('password', '').strip()
        if new_password:
            user.set_password(new_password)

        user.save()
        messages.success(request, f'Đã cập nhật user "{user.username}" thành công!')
        return redirect('admin_users')

    return render(request, 'admin/admin_user_form.html', {
        'action': 'edit',
        'user': user,
        'roles': User.Role.choices,
    })


@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_user_delete(request, user_id):
    user = get_object_or_404(User, id=user_id)

    if user == request.user:
        messages.error(request, 'Không thể xóa tài khoản đang đăng nhập!')
        return redirect('admin_users')

    if request.method == 'POST':
        username = user.username
        user.delete()
        messages.success(request, f'Đã xóa user "{username}" thành công!')
        return redirect('admin_users')

    return render(request, 'admin/admin_user_confirm_delete.html', {'user': user})

# ==================== CHI TIẾT ĐƠN HÀNG ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def admin_order_detail(request, order_id):
    """
    Trang chi tiết đơn hàng
    """

    orders = [
        {
            'id': 1234,
            'customer': 'Nguyễn Văn A',
            'phone': '0901234567',
            'product': 'VinFast VF8',
            'quantity': 1,
            'total': '1,200,000,000 VNĐ',
            'payment': 'Chuyển khoản',
            'status': 'pending',
            'status_text': 'Chờ xác nhận',
            'date': '01/02/2026',
            'address': 'TP.HCM'
        },
        {
            'id': 1233,
            'customer': 'Trần Thị B',
            'phone': '0912345678',
            'product': 'Tesla Model 3',
            'quantity': 1,
            'total': '1,500,000,000 VNĐ',
            'payment': 'Trả góp',
            'status': 'confirmed',
            'status_text': 'Đã xác nhận',
            'date': '31/01/2026',
            'address': 'Hà Nội'
        },
    ]

    # Tìm đơn hàng theo ID
    order = next((o for o in orders if o['id'] == order_id), None)
    if not order:
        messages.error(request, f'Không tìm thấy đơn hàng #{order_id}')
        return redirect('admin_donhang')

    return render(request, 'admin/admin_order_detail.html', {'order': order})
# ==================== API STATIONS ====================
from django.http import JsonResponse

def api_stations(request):
    from myapp.models import ChargingStation
    
    stations = ChargingStation.objects.all()
    
    data = []
    for s in stations:
        data.append({
            'id': s.id,
            'name': s.name,
            'address': s.address,
            'lat': s.latitude,
            'lon': s.longitude,
            'type': s.charger_type,
            'power': s.power,
            'total_ports': s.total_ports,
            'available_ports': s.available_ports,
            'status': s.status,
            'image': s.image.url if s.image else None,
        })
    
    return JsonResponse({'stations': data})

# ==================== SEARCH PAGE ====================
from django.views.decorators.csrf import csrf_exempt
import json

def search_page(request):
    from myapp.models import SearchHistory
    history = []
    if request.user.is_authenticated:
        history = SearchHistory.objects.filter(
            user=request.user
        ).values(
            'id', 'query', 'display_name',
            'latitude', 'longitude',
            'image_url', 'searched_at'
        )[:20]
    return render(request, 'search/search_page.html', {
        'history': list(history)
    })


@csrf_exempt
def api_search_history(request):
    from myapp.models import SearchHistory

    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            SearchHistory.objects.create(
                user=request.user if request.user.is_authenticated else None,
                query=data.get('query', ''),
                display_name=data.get('display_name', ''),
                latitude=data.get('latitude'),
                longitude=data.get('longitude'),
                image_url=data.get('image_url', ''),
            )
            return JsonResponse({'status': 'ok'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})

    elif request.method == 'GET':
        if request.user.is_authenticated:
            history = SearchHistory.objects.filter(
                user=request.user
            )
        else:
            history = SearchHistory.objects.filter(user=None)

        history = history.values(
            'id', 'query', 'display_name',
            'latitude', 'longitude',
            'image_url', 'searched_at'
        )[:20]

        history_list = []
        for item in history:
            item['searched_at'] = item['searched_at'].strftime('%Y-%m-%dT%H:%M:%S')
            history_list.append(item)

        return JsonResponse({'history': history_list})

    return JsonResponse({'status': 'error'}, status=405)

# ==================== export_orders_excel ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def export_orders_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo đơn hàng"

    # Định dạng tiêu đề (Header) cho đẹp
    header_font = Font(bold=True, color="FFFFFF")
    # Sử dụng PatternFill trực tiếp từ thư viện đã import
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    headers = ['Mã ĐH', 'Khách hàng', 'Sản phẩm', 'Tổng tiền (VNĐ)', 'Trạng thái', 'Ngày đặt']
    ws.append(headers)

    # Apply style cho header
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Lấy dữ liệu
    orders = Order.objects.all().select_related('user').prefetch_related('items__product').order_by('-created_at')

    for order in orders:
        product_names = ", ".join([item.product.name for item in order.items.all()])
        customer_name = f"{order.user.last_name} {order.user.first_name}" if order.user.last_name else order.user.username
        
        # Thêm dòng dữ liệu
        row = [
            f"#{order.id}",
            customer_name,
            product_names,
            float(order.total_price), 
            order.get_status_display(), 
            order.created_at.strftime('%d/%m/%Y')
        ]
        ws.append(row)

    # --- TỐI ƯU HIỂN THỊ ---

    # 1. Định dạng cột Tiền (Cột D) có dấu phân cách phần nghìn (Sửa lỗi 8.5E+08)
    for cell in ws['D']:
        if cell.row > 1: # Bỏ qua tiêu đề
            cell.number_format = '#,##0' 
            cell.alignment = Alignment(horizontal="right")

    # 2. Tự động giãn độ rộng cột (Auto-fit - Sửa lỗi mất chữ)
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter 
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 5

    # Xuất file
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=danh_sach_don_hang.xlsx'
    wb.save(response)
    return response

# ==================== export Nhập và xuất kho ====================
@login_required
@user_passes_test(admin_required, login_url='/login/')
def export_stock_in(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo nhập kho"

    # Style Header giống file đơn hàng của Huy
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    headers = ['ID', 'Sản phẩm', 'Số lượng', 'Giá nhập (VNĐ)', 'Người nhập', 'Ngày nhập']
    ws.append(headers)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Lấy dữ liệu
    queryset = StockIn.objects.select_related('product', 'imported_by').all().order_by('-imported_at')

    for item in queryset:
        ws.append([
            f"IN#{item.id}",
            item.product.name,
            item.quantity,
            float(item.import_price),
            item.imported_by.username,
            item.imported_at.strftime('%d/%m/%Y %H:%M')
        ])

    # Định dạng tiền và Auto-fit
    for cell in ws['D']:
        if cell.row > 1:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column].width = max_length + 5

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=bao_cao_nhap_kho.xlsx'
    wb.save(response)
    return response

# Hàm xuất file XUẤT KHO
@login_required
@user_passes_test(admin_required, login_url='/login/')
def export_stock_out(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo xuất kho"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    
    headers = ['ID', 'Sản phẩm', 'Số lượng', 'Người xuất', 'Ngày xuất']
    ws.append(headers)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    queryset = StockOut.objects.select_related('product', 'exported_by').all().order_by('-exported_at')

    for item in queryset:
        ws.append([
            f"OUT#{item.id}",
            item.product.name,
            item.quantity,
            item.exported_by.username,
            item.exported_at.strftime('%d/%m/%Y %H:%M')
        ])

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column].width = max_length + 5

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=bao_cao_xuat_kho.xlsx'
    wb.save(response)
    return response